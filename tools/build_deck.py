"""강의 PPT(docs/source/Campaign_3-2_Lecture_with_Guide.pptx) → public/deck.js
슬라이드마다 본문 줄 · 표 · 발표자 노트([멘트]/[진행]/[주의]…) · 하이퍼링크 · 'NN쪽' 참조를 뽑습니다.
슬라이드 그림(public/ppt/NNN.jpg)은 README의 방법으로 따로 만듭니다.
실행: python3 tools/build_deck.py
"""
import json, re, pathlib
from pptx import Presentation

ROOT = pathlib.Path(__file__).resolve().parent.parent
prs = Presentation(ROOT / 'docs/source/Campaign_3-2_Lecture_with_Guide.pptx')
HEAD = '3-2. 캠페인 기획 실습'

def walk(shapes):
    for sh in shapes:
        if sh.shape_type == 6:
            yield from walk(sh.shapes)
        else:
            yield sh

def links_of(sh):
    out = []
    try:
        a = sh.click_action.hyperlink.address
        if a: out.append((sh.text_frame.text.strip() if sh.has_text_frame else '', a))
    except Exception:
        pass
    frames = []
    if sh.has_text_frame: frames.append(sh.text_frame)
    if getattr(sh, 'has_table', False) and sh.has_table:
        frames += [c.text_frame for r in sh.table.rows for c in r.cells]
    for tf in frames:
        for p in tf.paragraphs:
            for r in p.runs:
                a = r.hyperlink.address
                if a: out.append((r.text.strip(), a))
    return out

slides = []
for i, s in enumerate(prs.slides, 1):
    lines, tables, links = [], [], []
    for sh in walk(s.shapes):
        links += links_of(sh)
        if getattr(sh, 'has_table', False) and sh.has_table:
            rows = []
            for r in sh.table.rows:
                cells = []
                for c in r.cells:
                    t = c.text.replace('\n', ' / ').strip()
                    if not cells or cells[-1] != t: cells.append(t)
                rows.append(cells)
            tables.append(rows)
        elif sh.has_text_frame and sh.text_frame.text.strip():
            lines.append(re.sub(r'\s*\n\s*', ' / ', sh.text_frame.text.strip()))
    notes = s.notes_slide.notes_text_frame.text.strip() if s.has_notes_slide else ''
    # 노트: "[태그] 내용" 단위로 나누기
    parts = [{'tag': m.group(1), 'text': m.group(2).strip()} for m in re.finditer(r'\[([^\]]{1,12})\]\s*(.*?)(?=\n?\[[^\]]{1,12}\]|\Z)', notes, re.S)]
    if notes and not parts: parts = [{'tag': '노트', 'text': notes}]
    body = [l for l in lines if l != HEAD]
    title = body[0] if body else ''
    step = body[1] if len(body) > 1 and len(body[1]) <= 22 and title.startswith('실습') else ''
    head = body[2] if step and len(body) > 2 else (body[1] if len(body) > 1 and not step else '')
    alltext = ' '.join(lines + [c for t in tables for r in t for c in r] + [notes])
    refs = sorted({int(n) for n in re.findall(r'(\d{1,3})\s*쪽', alltext) if 1 <= int(n) <= len(prs.slides) and int(n) != i})
    seen, uniq = set(), []
    for t, a in links:
        if a not in seen: seen.add(a); uniq.append({'t': t or a, 'u': a})
    slides.append({'n': i, 'hidden': s._element.get('show') == '0', 'title': title, 'step': step, 'head': head,
                   'lines': body, 'tables': tables, 'notes': parts, 'links': uniq, 'refs': refs})

js = '/* 자동 생성: python3 tools/build_deck.py — 원본 docs/source/Campaign_3-2_Lecture_with_Guide.pptx */\nconst DECK = ' + json.dumps(slides, ensure_ascii=False) + ';\n'
(ROOT / 'public/deck.js').write_text(js, encoding='utf-8')
print('slides', len(slides), 'with links', sum(1 for x in slides if x['links']), 'links', sum(len(x['links']) for x in slides), 'with refs', sum(1 for x in slides if x['refs']), 'KB', len(js.encode()) // 1024)
